from __future__ import annotations

import logging
import re
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from pypdf import PdfReader

from .config import Settings, ensure_runtime_dirs
from .db import connect, initialize_database

EXCEL_FILENAME = "采购合同细节.xlsx"
PDF_DIRNAME = "合同文件"
HEADER_ROW = 3
DATA_START_ROW = 4
LOGGER = logging.getLogger("pypdf")
LOGGER.setLevel(logging.ERROR)

REQUIRED_HEADERS = [
    "序号",
    "合同编号",
    "合同名称",
    "项目编号",
    "项目名称",
    "采购人(甲方)",
    "供应商名称(乙方)",
    "主要标的名称",
    "规格型号(或服务要求)",
    "主要标的数量",
    "主要标的单价(万元)",
    "合同金额（万元）",
    "采购方式",
    "合同签订日期",
    "网址",
    "合同文件",
]


@dataclass(frozen=True)
class PdfRecord:
    path: Path
    serial_no: int | None
    codes: list[str]
    size_bytes: int
    page_count: int | None
    text_status: str


@dataclass(frozen=True)
class ImportReport:
    source_excel: str
    dry_run: bool
    total_rows: int
    imported_rows: int
    matched_files: int
    unmatched_rows: int
    unmatched_files: int
    code_matched_files: int
    serial_prefix_matched_files: int
    warnings: list[str]
    errors: list[str]
    requires_review: list[str]
    blocked: bool
    allow_warnings: bool
    unmatched_row_serials: list[int]
    unmatched_file_names: list[str]
    db_path: str

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def extract_codes(value: object) -> list[str]:
    text = str(value or "")
    return re.findall(r"(?:ZCY)?\d{4,}", text, flags=re.IGNORECASE)


def extract_serial_prefix(path: Path) -> int | None:
    match = re.match(r"^\s*(\d+)\s*[.．]", path.name)
    return int(match.group(1)) if match else None


def to_int(value: object) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(str(value).strip()))
    except ValueError:
        return None


def to_float(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(str(value).strip())
    except ValueError:
        return None


def to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def to_iso_date(value: object) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).date().isoformat()
        except Exception:
            return ""
    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def classify_pdf(path: Path) -> tuple[int | None, str]:
    try:
        reader = PdfReader(str(path))
        page_count = len(reader.pages)
        text = ""
        for page in reader.pages[:3]:
            text += page.extract_text() or ""
        return page_count, "text_ok" if len(text.strip()) >= 50 else "low_or_scanned"
    except Exception:
        return None, "error"


def scan_pdfs(settings: Settings) -> list[PdfRecord]:
    pdf_dir = settings.raw_dir / PDF_DIRNAME
    records = []
    if not pdf_dir.exists():
        return records
    for path in sorted(pdf_dir.glob("*.pdf"), key=lambda item: (extract_serial_prefix(item) or 10**9, item.name)):
        page_count, text_status = classify_pdf(path)
        records.append(
            PdfRecord(
                path=path,
                serial_no=extract_serial_prefix(path),
                codes=extract_codes(path.stem),
                size_bytes=path.stat().st_size,
                page_count=page_count,
                text_status=text_status,
            )
        )
    return records


def read_excel_rows(settings: Settings) -> list[dict[str, object]]:
    excel_path = settings.raw_dir / EXCEL_FILENAME
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")
    workbook = load_workbook(excel_path, data_only=True, read_only=True)
    sheet = workbook["Sheet1"]
    headers = [to_text(cell.value) for cell in sheet[HEADER_ROW]]
    if headers[: len(REQUIRED_HEADERS)] != REQUIRED_HEADERS:
        raise ValueError("Excel headers do not match expected row 3 structure")
    rows = []
    for row in sheet.iter_rows(min_row=DATA_START_ROW, values_only=True):
        data = dict(zip(REQUIRED_HEADERS, row[: len(REQUIRED_HEADERS)]))
        if all(value is None or value == "" for value in data.values()):
            continue
        serial_no = to_int(data.get("序号"))
        if serial_no is None:
            continue
        rows.append(data)
    return rows


def normalize_contract(row: dict[str, object]) -> dict[str, object]:
    return {
        "serial_no": to_int(row.get("序号")),
        "contract_no": to_text(row.get("合同编号")),
        "contract_name": to_text(row.get("合同名称")),
        "project_no": to_text(row.get("项目编号")),
        "project_name": to_text(row.get("项目名称")),
        "purchaser": to_text(row.get("采购人(甲方)")),
        "supplier": to_text(row.get("供应商名称(乙方)")),
        "item_name": to_text(row.get("主要标的名称")),
        "spec_or_service": to_text(row.get("规格型号(或服务要求)")),
        "item_quantity": to_text(row.get("主要标的数量")),
        "unit_price_wan": to_float(row.get("主要标的单价(万元)")),
        "amount_wan": to_float(row.get("合同金额（万元）")),
        "purchase_method": to_text(row.get("采购方式")),
        "signed_date": to_iso_date(row.get("合同签订日期")),
        "source_url": to_text(row.get("网址")),
        "excel_contract_file": to_text(row.get("合同文件")),
    }


def import_excel(settings: Settings, dry_run: bool = False, allow_warnings: bool = False) -> ImportReport:
    ensure_runtime_dirs(settings)
    initialize_database(settings.db_path)
    excel_path = settings.raw_dir / EXCEL_FILENAME
    rows = read_excel_rows(settings)
    pdfs = scan_pdfs(settings)
    pdfs_by_serial = {record.serial_no: record for record in pdfs if record.serial_no is not None}
    pdfs_by_position = {index + 1: record for index, record in enumerate(pdfs)}
    matched_pdf_paths: set[Path] = set()
    unmatched_row_serials: list[int] = []
    warnings: list[str] = []
    errors: list[str] = []
    requires_review: list[str] = []
    prepared: list[tuple[dict[str, object], PdfRecord | None, str | None]] = []
    code_matched = 0
    serial_prefix_matched = 0

    serial_counts = Counter(normalize_contract(row)["serial_no"] for row in rows)
    contract_no_counts = Counter(normalize_contract(row)["contract_no"] for row in rows if normalize_contract(row)["contract_no"])
    for serial_no, count in sorted(serial_counts.items()):
        if serial_no is not None and count > 1:
            requires_review.append(f"serial {serial_no}: duplicate Excel serial appears {count} times")
    for contract_no, count in sorted(contract_no_counts.items()):
        if count > 1:
            requires_review.append(f"contract_no {contract_no}: duplicate Excel contract number appears {count} times")

    for row_index, row in enumerate(rows, start=1):
        contract = normalize_contract(row)
        serial_no = contract["serial_no"]
        pdf = pdfs_by_serial.get(serial_no)
        match_method = None
        if pdf is not None and pdf.path in matched_pdf_paths:
            requires_review.append(f"serial {serial_no}: duplicate Excel serial would reuse PDF {pdf.path.name}")
            pdf = None
        if pdf is None:
            if serial_no is not None:
                unmatched_row_serials.append(serial_no)
            prepared.append((contract, None, None))
            continue
        excel_codes = {code.upper() for code in extract_codes(contract["excel_contract_file"])}
        pdf_codes = {code.upper() for code in pdf.codes}
        if excel_codes and pdf_codes:
            if excel_codes & pdf_codes:
                match_method = "code"
                code_matched += 1
            else:
                requires_review.append(f"serial {serial_no}: Excel/PDF code mismatch")
                match_method = "serial_prefix"
                serial_prefix_matched += 1
        else:
            warnings.append(f"serial {serial_no}: matched PDF by serial prefix without contract code evidence")
            match_method = "serial_prefix"
            serial_prefix_matched += 1
        matched_pdf_paths.add(pdf.path)
        prepared.append((contract, pdf, match_method))

    unmatched_file_names = [record.path.name for record in pdfs if record.path not in matched_pdf_paths]
    blocked = bool(not dry_run and requires_review and not allow_warnings)

    if not dry_run and not blocked:
        with connect(settings.db_path) as conn:
            conn.execute("DELETE FROM contract_files")
            conn.execute("DELETE FROM contracts")
            conn.execute("DELETE FROM sqlite_sequence WHERE name IN ('contracts', 'contract_files')")
            for contract, pdf, match_method in prepared:
                cursor = conn.execute(
                    """
                    INSERT INTO contracts (
                        serial_no, contract_no, contract_name, project_no, project_name,
                        purchaser, supplier, item_name, spec_or_service, item_quantity,
                        unit_price_wan, amount_wan, purchase_method, signed_date, source_url,
                        excel_contract_file
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        contract["serial_no"],
                        contract["contract_no"],
                        contract["contract_name"],
                        contract["project_no"],
                        contract["project_name"],
                        contract["purchaser"],
                        contract["supplier"],
                        contract["item_name"],
                        contract["spec_or_service"],
                        contract["item_quantity"],
                        contract["unit_price_wan"],
                        contract["amount_wan"],
                        contract["purchase_method"],
                        contract["signed_date"],
                        contract["source_url"],
                        contract["excel_contract_file"],
                    ),
                )
                if pdf is not None and match_method is not None:
                    conn.execute(
                        """
                        INSERT INTO contract_files (
                            contract_id, original_filename, relative_path, file_ext, size_bytes,
                            page_count, text_status, match_method
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            cursor.lastrowid,
                            pdf.path.name,
                            pdf.path.relative_to(settings.raw_dir).as_posix(),
                            pdf.path.suffix.lower(),
                            pdf.size_bytes,
                            pdf.page_count,
                            pdf.text_status,
                            match_method,
                        ),
                    )
            conn.execute(
                """
                INSERT INTO import_runs (
                    source_excel, total_rows, imported_rows, matched_files, unmatched_rows,
                    unmatched_files, code_matched_files, serial_prefix_matched_files,
                    finished_at, status, message
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP, ?, ?)
                """,
                (
                    excel_path.relative_to(settings.raw_dir).as_posix(),
                    len(rows),
                    len(prepared),
                    len(matched_pdf_paths),
                    len(unmatched_row_serials),
                    len(unmatched_file_names),
                    code_matched,
                    serial_prefix_matched,
                    "success" if not (warnings or requires_review or errors) else "success_with_warnings",
                    "; ".join((errors + requires_review + warnings)[:10]),
                ),
            )

    return ImportReport(
        source_excel=excel_path.relative_to(settings.raw_dir).as_posix(),
        dry_run=dry_run,
        total_rows=len(rows),
        imported_rows=len(prepared),
        matched_files=len(matched_pdf_paths),
        unmatched_rows=len(unmatched_row_serials),
        unmatched_files=len(unmatched_file_names),
        code_matched_files=code_matched,
        serial_prefix_matched_files=serial_prefix_matched,
        warnings=warnings,
        errors=errors,
        requires_review=requires_review,
        blocked=blocked,
        allow_warnings=allow_warnings,
        unmatched_row_serials=unmatched_row_serials,
        unmatched_file_names=unmatched_file_names,
        db_path=str(settings.db_path),
    )


def check_files(settings: Settings) -> dict[str, object]:
    initialize_database(settings.db_path)
    missing = []
    with connect(settings.db_path) as conn:
        rows = conn.execute("SELECT id, relative_path FROM contract_files ORDER BY id").fetchall()
    for row in rows:
        path = (settings.raw_dir / row["relative_path"]).resolve()
        try:
            path.relative_to(settings.raw_dir)
        except ValueError:
            missing.append({"file_id": row["id"], "relative_path": row["relative_path"], "reason": "outside_raw_dir"})
            continue
        if not path.exists():
            missing.append({"file_id": row["id"], "relative_path": row["relative_path"], "reason": "missing"})
    return {"checked": len(rows), "missing": missing, "missing_count": len(missing)}
